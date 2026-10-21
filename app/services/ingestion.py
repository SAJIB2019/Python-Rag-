import io
import json
import mimetypes
import re
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from llama_index.core.node_parser import SentenceSplitter
from openpyxl import load_workbook
from pypdf import PdfReader


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\t", " ").replace("\x00", "")
    return re.sub(r"\n{3,}", "\n\n", re.sub(r" {2,}", " ", text)).strip()


def chunk_text(text: str, size: int = 200, overlap: int = 30) -> list[dict]:
    if size <= 0 or overlap < 0 or overlap >= size:
        raise ValueError("Chunk size must be positive and overlap smaller than size")
    if not text.strip():
        return []
    from llama_index.core.utils import get_tokenizer

    tokenizer = get_tokenizer()
    splitter = SentenceSplitter(chunk_size=size, chunk_overlap=overlap, tokenizer=tokenizer)
    parts = [part.strip() for part in splitter.split_text(text) if part.strip()]
    return [
        {"chunk_index": i, "content": part, "token_count": len(tokenizer(part))}
        for i, part in enumerate(parts)
    ]


MIME_KINDS = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "text/csv": "csv",
    "text/plain": "txt",
    "text/markdown": "txt",
    "application/json": "json",
    "text/html": "html",
}
EXT_KINDS = {
    "pdf": "pdf",
    "docx": "docx",
    "xlsx": "xlsx",
    "csv": "csv",
    "txt": "txt",
    "md": "txt",
    "markdown": "txt",
    "json": "json",
    "html": "html",
    "htm": "html",
}


def extract_file(data: bytes, filename: str, mime_type: str) -> tuple[str, str]:
    if not data:
        raise ValueError("Uploaded file is empty")
    kind = MIME_KINDS.get(mime_type) or EXT_KINDS.get(Path(filename).suffix.lower()[1:])
    if kind is None:
        raise ValueError("Supported uploads: PDF, DOCX, XLSX, CSV, TXT, Markdown, JSON, HTML")
    if kind == "pdf":
        text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
    elif kind == "docx":
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        parts.extend("\t".join(c.text for c in row.cells) for t in doc.tables for row in t.rows)
        text = "\n".join(parts)
    elif kind == "xlsx":
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            text = "\n\n".join(
                f"Sheet: {sheet.title}\n"
                + "\n".join(
                    "\t".join(str(c) if c is not None else "" for c in row)
                    for row in sheet.iter_rows(values_only=True)
                )
                for sheet in book
            )
        finally:
            book.close()
    else:
        text = data.decode("utf-8-sig", errors="replace")
        if kind == "json":
            try:
                text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
            except ValueError:
                pass
        elif kind == "html":
            soup = BeautifulSoup(text, "html.parser")
            for tag in soup(["script", "style"]):
                tag.decompose()
            text = soup.get_text(" ", strip=True)
    text = text.replace("\x00", "").strip()
    if not text:
        raise ValueError("No readable text could be extracted from the uploaded file")
    return text, mime_type or mimetypes.guess_type(filename)[0] or "application/octet-stream"
