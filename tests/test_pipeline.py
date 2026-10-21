import io
import json
from unittest.mock import AsyncMock

import httpx
import pytest
from docx import Document
from openpyxl import Workbook
from pydantic import ValidationError

from app.core.config import Settings
from app.integrations.providers import (
    FakeEmbeddingProvider,
    FakeLlmProvider,
    OpenAICompatibleProvider,
)
from app.schemas import CreateDocument, QueryInput, RetrievedChunk
from app.services.ingestion import chunk_text, clean_text, extract_file
from app.services.retrieval import QueryService, build_context, merge_hybrid_results


def chunk(index=0, content="Refunds are available within 30 days."):
    return RetrievedChunk(
        documentId="doc", filename="policy.txt", chunkIndex=index, content=content, score=0.1
    )


def test_rrf_merges_and_ranks_shared_results():
    results = merge_hybrid_results([chunk(0), chunk(1)], [chunk(1), chunk(2)])
    assert [r.chunkIndex for r in results] == [1, 0, 2]
    assert results[0].sources == ["vector", "keyword"]
    assert results[0].hybridScore == pytest.approx(1 / 62 + 1 / 61)


def test_context_preserves_first_chunk_and_citations():
    results = merge_hybrid_results([chunk(0, "a" * 400), chunk(1)], [])
    context = build_context(results, max_characters=100)
    assert context.text == "[source:1] " + "a" * 400
    assert len(context.sources) == 1
    assert context.sources[0].snippet == "a" * 240 + "..."
    assert build_context([]).model_dump() == {"text": "", "sources": []}


def test_chunking_and_cleaning():
    assert clean_text("  a\t  b\r\n\n\n\x00c  ") == "a b\n\nc"
    chunks = chunk_text("Refunds are available within thirty days. " * 80, 40, 5)
    assert len(chunks) > 1
    assert all(0 < c["token_count"] <= 40 for c in chunks)
    assert [c["chunk_index"] for c in chunks] == list(range(len(chunks)))
    assert chunk_text("   ") == []
    with pytest.raises(ValueError, match="overlap"):
        chunk_text("text", 20, 20)


@pytest.mark.parametrize(
    "body",
    [{"question": "  "}, {"question": "x" * 2001}, {"question": 2}, {"question": "q", "extra": 1}],
)
def test_query_validation(body):
    with pytest.raises(ValidationError):
        QueryInput.model_validate(body)


def test_document_validation_and_trim():
    model = CreateDocument(filename=" policy.txt ", mimeType="text/plain", content="  hello  ")
    assert (model.filename, model.content) == ("policy.txt", "hello")
    with pytest.raises(ValidationError):
        CreateDocument(filename="a", mimeType="text/html", content="x")


@pytest.mark.parametrize(
    ("data", "name", "mime", "expected"),
    [
        (b"hello", "a.txt", "text/plain", "hello"),
        (b"# Title", "a.md", "application/octet-stream", "# Title"),
        (b'{"a":1}', "a.json", "application/json", '"a": 1'),
        (b"invalid json", "a.json", "application/json", "invalid json"),
        (
            b"<script>bad()</script><style>bad</style><p>Hello &amp; world</p>",
            "a.html",
            "text/html",
            "Hello & world",
        ),
        (b"name,value\na,1", "a.csv", "text/csv", "a,1"),
    ],
)
def test_text_uploads(data, name, mime, expected):
    assert expected in extract_file(data, name, mime)[0]


def test_office_uploads():
    document = Document()
    document.add_paragraph("Refund policy")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "30 days"
    data = io.BytesIO()
    document.save(data)
    text, _ = extract_file(data.getvalue(), "a.docx", "")
    assert "Refund policy" in text and "30 days" in text
    workbook = Workbook()
    workbook.active.append(["Refund", 30])
    data = io.BytesIO()
    workbook.save(data)
    assert "Refund\t30" in extract_file(data.getvalue(), "a.xlsx", "")[0]


def test_pdf_upload():
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 10 200 Td (Refunds within 30 days) Tj ET")
    page[NameObject("/Contents")] = stream
    buffer = io.BytesIO()
    writer.write(buffer)
    assert (
        "Refunds within 30 days"
        in extract_file(buffer.getvalue(), "policy.pdf", "application/pdf")[0]
    )


@pytest.mark.parametrize(("data", "name"), [(b"", "a.txt"), (b"  ", "a.txt"), (b"x", "a.exe")])
def test_bad_uploads(data, name):
    with pytest.raises(ValueError):
        extract_file(data, name, "application/octet-stream")


async def test_fake_providers_and_query():
    vector = await FakeEmbeddingProvider(8).embed("hello")
    assert sum(v * v for v in vector) == pytest.approx(1, abs=0.00001)
    search = AsyncMock()
    search.vector_search.return_value = [chunk()]
    search.keyword_search.return_value = [chunk()]
    response = await QueryService(search, FakeEmbeddingProvider(8), FakeLlmProvider()).query(
        "Refund?"
    )
    assert response.sources[0].filename == "policy.txt"
    assert response.debug.results[0].sources == ["vector", "keyword"]
    search.vector_search.return_value = []
    search.keyword_search.return_value = []
    response = await QueryService(search, FakeEmbeddingProvider(8), FakeLlmProvider()).query(
        "Refund?"
    )
    assert "don't have enough" in response.answer and response.sources == []


async def test_remote_provider_contract():
    requests = []

    def respond(request):
        requests.append(request)
        if request.url.path.endswith("embeddings"):
            return httpx.Response(200, json={"data": [{"embedding": [0.1, 0.2]}]})
        return httpx.Response(
            200, json={"choices": [{"message": {"content": "30 days [source:1]"}}]}
        )

    settings = Settings(
        _env_file=None,
        embedding_dimensions=2,
        openai_compatible_base_url="https://provider.test/v1/",
        openai_compatible_api_key="secret",
    )
    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        provider = OpenAICompatibleProvider(settings, client)
        assert await provider.embed("text") == [0.1, 0.2]
        assert await provider.generate("Refund?", "[source:1] 30 days") == "30 days [source:1]"
    assert requests[0].headers["authorization"] == "Bearer secret"
    assert json.loads(requests[0].content)["dimensions"] == 2
    assert "[source:1] 30 days" in json.loads(requests[1].content)["messages"][1]["content"]


@pytest.mark.parametrize("vector", [[1], [1, "bad"], [True, 1], [None, 1]])
async def test_remote_rejects_invalid_vectors(vector):
    settings = Settings(
        _env_file=None,
        embedding_dimensions=2,
        openai_compatible_base_url="https://provider.test/v1",
    )
    async with httpx.AsyncClient(
        transport=httpx.MockTransport(
            lambda r: httpx.Response(200, json={"data": [{"embedding": vector}]})
        )
    ) as client:
        with pytest.raises(ValueError):
            await OpenAICompatibleProvider(settings, client).embed("text")


def test_production_config_requires_real_providers():
    with pytest.raises(ValidationError, match="fake"):
        Settings(_env_file=None, node_env="production")
